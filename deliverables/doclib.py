"""Tiny document model rendered two ways: DOCX (python-docx, editable) and PDF (HTML printed by headless Chrome).

Both renderings use A4, the same margins, Arial and the same font sizes, so the PDF page count is a good guide
to the DOCX page count.
"""
from __future__ import annotations

import base64
import html
import re
import subprocess
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document as DocxDocument
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
ACCENT = "1F3A68"
HEADER_FILL = "DCE4F2"
INLINE = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)")


# --------------------------------------------------------------------------- blocks
@dataclass
class H:
    level: int
    text: str


@dataclass
class P:
    text: str
    size: float | None = None
    align: str | None = None          # "center" | "right"
    italic: bool = False
    space_after: float = 4
    keep_with_next: bool = False


@dataclass
class Bullets:
    items: list[str]
    numbered: bool = False
    size: float | None = None


@dataclass
class Table:
    header: list[str]
    rows: list[list[str]]
    widths_cm: list[float]
    size: float | None = None
    caption: str | None = None


@dataclass
class Img:
    path: Path
    width_cm: float
    caption: str | None = None


@dataclass
class PageBreak:
    pass


@dataclass
class Doc:
    title: str
    blocks: list
    subtitle: str | None = None
    meta: list[str] = field(default_factory=list)
    font_size: float = 10.5
    margins_cm: float = 2.0
    footer: str | None = None


# --------------------------------------------------------------------------- DOCX
def _runs(paragraph, text: str, size: float, italic: bool = False, bold: bool = False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run, b, i, mono = paragraph.add_run(part[2:-2]), True, italic, False
        elif part.startswith("`") and part.endswith("`"):
            run, b, i, mono = paragraph.add_run(part[1:-1]), bold, italic, True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run, b, i, mono = paragraph.add_run(part[1:-1]), bold, True, False
        else:
            run, b, i, mono = paragraph.add_run(part), bold, italic, False
        run.bold, run.italic = b or None, i or None
        run.font.size = Pt(size * (0.92 if mono else 1))
        if mono:
            run.font.name = "Courier New"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")


def _shade(cell, fill: str):
    tc_pr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _page_number_field(paragraph, size):
    run = paragraph.add_run()
    run.font.size = Pt(size)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._element.append(el)


def to_docx(doc: Doc, path: Path) -> None:
    d = DocxDocument()
    section = d.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(doc.margins_cm))

    normal = d.styles["Normal"]
    normal.font.name = "Arial"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(doc.font_size)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.08
    for level, size in ((1, doc.font_size + 3.5), (2, doc.font_size + 1.5), (3, doc.font_size + 0.5)):
        style = d.styles[f"Heading {level}"]
        style.font.name, style.font.size, style.font.bold = "Arial", Pt(size), True
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.color.rgb = RGBColor.from_string(ACCENT)
        style.paragraph_format.space_before = Pt(10 if level == 1 else 7)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.keep_with_next = True

    title = d.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc.title)
    run.bold, run.font.size = True, Pt(doc.font_size + 7)
    run.font.color.rgb = RGBColor.from_string(ACCENT)
    title.paragraph_format.space_after = Pt(2)
    if doc.subtitle:
        sub = d.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _runs(sub, doc.subtitle, doc.font_size + 1.5, italic=False)
        sub.paragraph_format.space_after = Pt(2)
    for line in doc.meta:
        meta = d.add_paragraph()
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _runs(meta, line, doc.font_size - 0.5)
        meta.paragraph_format.space_after = Pt(1)
    if doc.subtitle or doc.meta:
        rule = d.add_paragraph()
        p_pr = rule._element.get_or_add_pPr()
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "1"), ("w:color", ACCENT)):
            bottom.set(qn(k), v)
        border.append(bottom)
        p_pr.append(border)
        rule.paragraph_format.space_after = Pt(4)

    for block in doc.blocks:
        if isinstance(block, H):
            d.add_heading(block.text, level=block.level)
        elif isinstance(block, P):
            para = d.add_paragraph()
            _runs(para, block.text, block.size or doc.font_size, italic=block.italic)
            para.paragraph_format.space_after = Pt(block.space_after)
            para.paragraph_format.keep_with_next = block.keep_with_next or None
            if block.align:
                para.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}[block.align]
        elif isinstance(block, Bullets):
            for item in block.items:
                para = d.add_paragraph(style="List Number" if block.numbered else "List Bullet")
                _runs(para, item, block.size or doc.font_size)
                para.paragraph_format.space_after = Pt(1.5)
        elif isinstance(block, Table):
            size = block.size or doc.font_size - 1
            table = d.add_table(rows=1 + len(block.rows), cols=len(block.header))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            for r, cells in enumerate([block.header] + block.rows):
                for c, text in enumerate(cells):
                    cell = table.cell(r, c)
                    cell.width = Cm(block.widths_cm[c])
                    para = cell.paragraphs[0]
                    para.paragraph_format.space_after = Pt(0)
                    para.paragraph_format.line_spacing = 1.0
                    _runs(para, text, size, bold=r == 0)
                    if r == 0:
                        _shade(cell, HEADER_FILL)
            if block.caption:
                cap = d.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _runs(cap, block.caption, doc.font_size - 1.5, italic=True)
            else:
                d.add_paragraph().paragraph_format.space_after = Pt(2)
        elif isinstance(block, Img):
            d.add_picture(str(block.path), width=Cm(block.width_cm))
            d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            d.paragraphs[-1].paragraph_format.space_after = Pt(2)
            if block.caption:
                cap = d.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _runs(cap, block.caption, doc.font_size - 1.5, italic=True)
        elif isinstance(block, PageBreak):
            d.add_page_break()

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if doc.footer:
        _runs(footer, doc.footer + "  ·  page ", doc.font_size - 2)
    _page_number_field(footer, doc.font_size - 2)
    d.core_properties.title = doc.title
    d.save(str(path))


# --------------------------------------------------------------------------- HTML / PDF
def _inline_html(text: str) -> str:
    out = []
    for part in INLINE.split(text):
        if part.startswith("**") and part.endswith("**"):
            out.append(f"<b>{html.escape(part[2:-2])}</b>")
        elif part.startswith("`") and part.endswith("`"):
            out.append(f"<code>{html.escape(part[1:-1])}</code>")
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            out.append(f"<i>{html.escape(part[1:-1])}</i>")
        else:
            out.append(html.escape(part))
    return "".join(out)


def to_html(doc: Doc) -> str:
    fs = doc.font_size
    css = f"""
    @page {{ size: A4; margin: {doc.margins_cm}cm; }}
    body {{ font-family: Arial, Helvetica, sans-serif; font-size: {fs}pt; line-height: 1.24; color: #111; margin: 0; }}
    h1, h2, h3 {{ color: #{ACCENT}; margin: 0; break-after: avoid; }}
    h1 {{ font-size: {fs + 3.5}pt; margin-top: 10pt; margin-bottom: 3pt; }}
    h2 {{ font-size: {fs + 1.5}pt; margin-top: 7pt; margin-bottom: 3pt; }}
    h3 {{ font-size: {fs + 0.5}pt; margin-top: 7pt; margin-bottom: 3pt; }}
    p {{ margin: 0 0 4pt 0; }}
    ul, ol {{ margin: 0 0 4pt 0; padding-left: 0.63cm; }}
    li {{ margin-bottom: 1.5pt; }}
    table {{ border-collapse: collapse; margin: 0 auto 3pt auto; font-size: {fs - 1}pt; line-height: 1.12;
             break-inside: auto; }}
    tr {{ break-inside: avoid; }}
    th, td {{ border: 0.75pt solid #000; padding: 1.5pt 4pt; vertical-align: top; text-align: left; }}
    th {{ background: #{HEADER_FILL}; }}
    code {{ font-family: 'Courier New', monospace; font-size: 0.92em; }}
    .title {{ text-align: center; font-size: {fs + 7}pt; font-weight: bold; color: #{ACCENT}; margin: 0 0 2pt 0; }}
    .subtitle {{ text-align: center; font-size: {fs + 1.5}pt; margin: 0 0 2pt 0; }}
    .meta {{ text-align: center; font-size: {fs - 0.5}pt; margin: 0 0 1pt 0; }}
    .rule {{ border-bottom: 1pt solid #{ACCENT}; margin: 0 0 6pt 0; height: 4pt; }}
    .caption {{ text-align: center; font-style: italic; font-size: {fs - 1.5}pt; margin: 0 0 5pt 0; }}
    figure {{ margin: 0; text-align: center; break-inside: avoid; }}
    .pb {{ break-after: page; }}
    """
    parts = [f"<div class='title'>{html.escape(doc.title)}</div>"]
    if doc.subtitle:
        parts.append(f"<div class='subtitle'>{_inline_html(doc.subtitle)}</div>")
    parts += [f"<div class='meta'>{_inline_html(m)}</div>" for m in doc.meta]
    if doc.subtitle or doc.meta:
        parts.append("<div class='rule'></div>")
    for block in doc.blocks:
        if isinstance(block, H):
            parts.append(f"<h{block.level}>{html.escape(block.text)}</h{block.level}>")
        elif isinstance(block, P):
            style = []
            if block.size:
                style.append(f"font-size:{block.size}pt")
            if block.align:
                style.append(f"text-align:{block.align}")
            if block.italic:
                style.append("font-style:italic")
            style.append(f"margin-bottom:{block.space_after}pt")
            parts.append(f"<p style='{';'.join(style)}'>{_inline_html(block.text)}</p>")
        elif isinstance(block, Bullets):
            tag = "ol" if block.numbered else "ul"
            size = f" style='font-size:{block.size}pt'" if block.size else ""
            parts.append(f"<{tag}{size}>" + "".join(f"<li>{_inline_html(i)}</li>" for i in block.items) + f"</{tag}>")
        elif isinstance(block, Table):
            size = f"font-size:{block.size}pt;" if block.size else ""
            cols = "".join(f"<col style='width:{w}cm'>" for w in block.widths_cm)
            head = "".join(f"<th>{_inline_html(h)}</th>" for h in block.header)
            body = "".join("<tr>" + "".join(f"<td>{_inline_html(c)}</td>" for c in row) + "</tr>" for row in block.rows)
            parts.append(f"<table style='{size}width:{sum(block.widths_cm)}cm'><colgroup>{cols}</colgroup>"
                         f"<thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>")
            if block.caption:
                parts.append(f"<div class='caption'>{_inline_html(block.caption)}</div>")
        elif isinstance(block, Img):
            data = base64.b64encode(Path(block.path).read_bytes()).decode()
            cap = f"<div class='caption'>{_inline_html(block.caption)}</div>" if block.caption else ""
            parts.append(f"<figure><img src='data:image/png;base64,{data}' style='width:{block.width_cm}cm'>{cap}</figure>")
        elif isinstance(block, PageBreak):
            parts.append("<div class='pb'></div>")
    return f"<!doctype html><html><head><meta charset='utf-8'><style>{css}</style></head><body>{''.join(parts)}</body></html>"


def to_pdf(doc: Doc, path: Path) -> int:
    page = path.with_suffix(".tmp.html")
    page.write_text(to_html(doc))
    subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                    f"--print-to-pdf={path}", page.as_uri()], check=True, capture_output=True, timeout=120)
    page.unlink()
    info = subprocess.run(["pdfinfo", str(path)], capture_output=True, text=True).stdout
    return int(re.search(r"Pages:\s+(\d+)", info).group(1))


def build(doc: Doc, stem: Path) -> int:
    to_docx(doc, stem.with_suffix(".docx"))
    pages = to_pdf(doc, stem.with_suffix(".pdf"))
    print(f"{stem.name}: {pages} page(s) -> .docx + .pdf")
    return pages
